import time
import openpyxl
from openpyxl.styles import PatternFill
from selenium.webdriver.support.select import Select

def getreaddata(file,Sheet,rownum,columnnum):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[Sheet]
    return sheet.cell(rownum,columnnum).value

def getwritedata(file,Sheet,rownum,columnnum,data):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[Sheet]
    sheet.cell(rownum,columnnum).value = data
    workbook.save(file)
def getrowcount(file,Sheet):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[Sheet]
    return(sheet.max_row)

def getcolumncount(file,Sheet):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[Sheet]
    return(sheet.max_column)

def Fillgreencolor(file,Sheet,rownum,colnum):
     workbook = openpyxl.load_workbook(file)
     sheet = workbook[Sheet]
     greenfill = PatternFill(start_color='60b212',
                             end_color='60b212',
                             fill_type='solid')
     sheet.cell(rownum,colnum).fill = greenfill
     workbook.save(file)

def Fillredcolor(file,Sheet,rownum,colnum):
     workbook = openpyxl.load_workbook(file)
     sheet = workbook[Sheet]
     redfill = PatternFill(start_color='ff0000',
                             end_color='ff0000',
                             fill_type='solid')
     sheet.cell(rownum,colnum).fill = redfill
     workbook.save(file)
