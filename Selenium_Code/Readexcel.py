import time
import openpyxl

file = "C:\\Users\\91766\\OneDrive\\Documents\\Selenium practice\\Readsample1.xlsx"
workbook=openpyxl.load_workbook(file)
sheet = workbook.active

rows = sheet.max_row
cols = sheet.max_column
print(rows, cols)
time.sleep(3)

for r in range(1,rows+1):
    for c in range(1,cols+1):
        print(sheet.cell(r,c).value, end="        ")
    print()

