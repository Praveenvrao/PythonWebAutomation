import time
import openpyxl

# file = "C:\\Users\\91766\\OneDrive\\Documents\\Selenium practice\\Writesample1.xlsx"
# workbook=openpyxl.load_workbook(file)
# sheet = workbook.active
#
# for r in range(1,5):
#     for c in range(1,4):
#         sheet.cell(r,c).value = "PYTHON"
# workbook.save(file)

#2
file = "C:\\Users\\91766\\OneDrive\\Documents\\Selenium practice\\Writesample1.xlsx"
workbook=openpyxl.load_workbook(file)
sheet = workbook["Sheet2"]

sheet.cell(1,1).value = "S no"
sheet.cell(1,2).value = "Language"
sheet.cell(1,3).value = "Fee"

sheet.cell(2,1).value = 1
sheet.cell(2,2).value = "JAVA"
sheet.cell(2,3).value = 2000

sheet.cell(3,1).value = 2
sheet.cell(3,2).value = "PYTHON"
sheet.cell(3,3).value = 3000

sheet.cell(4,1).value = 3
sheet.cell(4,2).value = "C+"
sheet.cell(4,3).value = 20000

workbook.save(file)


